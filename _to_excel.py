"""Build Operations-Reporting-Requirements.xlsx from the markdown pack."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
DOCS = ROOT / "docs"
OUT = ROOT / "Operations-Reporting-Requirements.xlsx"

NAVY = "1F4E79"
GOLD = "C4A35A"
LIGHT = "E8EEF4"
WHITE = "FFFFFF"
INK = "1A1A1A"
MUTED = "5C5C5C"

fill_navy = PatternFill("solid", fgColor=NAVY)
fill_gold = PatternFill("solid", fgColor=GOLD)
fill_light = PatternFill("solid", fgColor=LIGHT)
fill_white = PatternFill("solid", fgColor=WHITE)
font_head = Font(name="Calibri", bold=True, color=WHITE, size=11)
font_title = Font(name="Calibri", bold=True, color=NAVY, size=18)
font_h2 = Font(name="Calibri", bold=True, color=NAVY, size=13)
font_body = Font(name="Calibri", color=INK, size=11)
font_small = Font(name="Calibri", color=MUTED, size=10)
thin = Border(
    left=Side(style="thin", color="C8C8C8"),
    right=Side(style="thin", color="C8C8C8"),
    top=Side(style="thin", color="C8C8C8"),
    bottom=Side(style="thin", color="C8C8C8"),
)
wrap = Alignment(wrap_text=True, vertical="top")


def split_row(line: str) -> list[str]:
    parts = [p.strip() for p in line.strip().strip("|").split("|")]
    return [re.sub(r"\*\*(.+?)\*\*", r"\1", p).replace("`", "") for p in parts]


def is_sep(line: str) -> bool:
    s = line.strip()
    return bool(s.startswith("|") and re.match(r"^\|[\s:|-]+\|?\s*$", s))


def iter_tables(path: Path):
    text = path.read_text(encoding="utf-8")
    heading = path.stem
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("#"):
            heading = line.lstrip("#").strip()
        if (
            line.strip().startswith("|")
            and i + 1 < len(lines)
            and is_sep(lines[i + 1])
        ):
            header = split_row(line)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|") and not is_sep(lines[i]):
                row = split_row(lines[i])
                while len(row) < len(header):
                    row.append("")
                rows.append(row[: len(header)])
                i += 1
            yield heading, header, rows
        else:
            i += 1


def parse_glossary(path: Path) -> list[list[str]]:
    text = path.read_text(encoding="utf-8")
    rows = []
    group = ""
    term = None
    definition = []
    avoid = ""
    def flush():
        nonlocal term, definition, avoid
        if term:
            rows.append([group, term, " ".join(definition).strip(), avoid])
        term, definition, avoid = None, [], ""

    for line in text.splitlines():
        if line.startswith("## ") and not line.startswith("## Language"):
            flush()
            group = line[3:].strip()
            continue
        m = re.match(r"^\*\*(.+?)\*\*:\s*$", line)
        if m:
            flush()
            term = m.group(1)
            continue
        if line.startswith("_Avoid_:"):
            avoid = line.split(":", 1)[1].strip()
            continue
        if term and line.strip() and not line.startswith("#"):
            definition.append(line.strip())
    flush()
    return rows


def parse_stories(path: Path) -> list[list[str]]:
    text = path.read_text(encoding="utf-8")
    epic = ""
    rows = []
    for line in text.splitlines():
        if line.startswith("## Epic"):
            epic = line.lstrip("#").strip()
        m = re.match(r"^(\d+)\.\s+As (.+?), I want (.+), so that (.+)\s*$", line)
        if m:
            rows.append([m.group(1), epic, "As " + m.group(2), m.group(3).rstrip("."), m.group(4).rstrip(".")])
    return rows


def parse_principles() -> list[list[str]]:
    return [
        ["R1", "Never mix Funded, Correspondent purchase, and Investor purchase in one unlabeled funded number."],
        ["R2", "Cycle time is p50 and p90, not average alone."],
        ["R3", "Pull-through is a named cohort (start month, lock month, or submission month)."],
        ["R4", "As-of Pipeline and event-dated Funded are different questions; label the time basis."],
        ["R5", "Waiting-on party is a first-class aging slice."],
        ["R6", "Initial LE send uses TRID general business days (creditor-open). CD wait uses TRID specific business days."],
        ["R7", "Missing events render unavailable, not zero."],
        ["R8", "Channel is required on every enterprise view."],
        ["R9", "Manufacturing QC is not independent audit QC."],
        ["R10", "LOS remains the system of action; BI may export stuck files only."],
    ]


def style_header(ws, ncols: int, row: int = 1) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row, col)
        cell.fill = fill_navy
        cell.font = font_head
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin
    ws.auto_filter.ref = f"A{row}:{get_column_letter(ncols)}{row}"
    ws.freeze_panes = f"A{row + 1}"
    ws.row_dimensions[row].height = 22


def write_table(ws, headers: list[str], rows: list[list[str]], start_row: int = 1) -> int:
    for c, h in enumerate(headers, 1):
        ws.cell(start_row, c, h)
    style_header(ws, len(headers), start_row)
    r = start_row + 1
    for row in rows:
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.font = font_body
            cell.alignment = wrap
            cell.border = thin
            if r % 2 == 0:
                cell.fill = fill_light
        ws.row_dimensions[r].height = min(90, 15 + 12 * max((str(x).count("\n") + len(str(x)) // 60) for x in row))
        r += 1
    widths = [len(h) for h in headers]
    for row in rows:
        for i, val in enumerate(row):
            widths[i] = max(widths[i], min(48, len(str(val))))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(42, max(12, w + 2))
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{r - 1}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return r - 1


def add_title(ws, title: str, subtitle: str = "") -> None:
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = title
    c.font = font_title
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws.merge_cells("A2:F2")
        d = ws["A2"]
        d.value = subtitle
        d.font = font_small
        ws.row_dimensions[2].height = 18


def collect_tables(path: Path, family_prefix: str | None = None) -> tuple[list[str], list[list[str]]]:
    """Union headers; prepend Family."""
    all_headers: list[str] = ["Family"]
    seen = {"family"}
    blocks = []
    for heading, header, rows in iter_tables(path):
        fam = heading
        for h in header:
            key = h.lower()
            if key not in seen:
                seen.add(key)
                all_headers.append(h)
        blocks.append((fam, header, rows))
    out_rows = []
    index = {h.lower(): i for i, h in enumerate(all_headers)}
    for fam, header, rows in blocks:
        for row in rows:
            rec = [""] * len(all_headers)
            rec[0] = fam
            for h, v in zip(header, row):
                rec[index[h.lower()]] = v
            out_rows.append(rec)
    return all_headers, out_rows


def cover(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Cover"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Operations reporting requirements"
    ws["A1"].font = Font(name="Calibri", bold=True, color=NAVY, size=22)
    ws.merge_cells("A1:B1")
    ws["A3"] = "Subject area"
    ws["B3"] = "Operations — loan manufacturing at a multi-channel non-bank mortgage company"
    ws["A4"] = "Channels"
    ws["B4"] = "Retail, Wholesale, Correspondent"
    ws["A5"] = "Lifecycle"
    ws["B5"] = "File start / intake through post-close, trailing documents, and investor delivery"
    ws["A6"] = "Companion markdown"
    ws["B6"] = "Operations-Reporting-Requirements.md"
    ws["A7"] = "Wireframes"
    ws["B7"] = "docs/wireframes/index.html (by department P01–P15 + EXE)"
    ws["A8"] = "Status"
    ws["B8"] = "Working requirements. ASSUMED SLA is not company policy."
    for r in range(3, 9):
        ws.cell(r, 1).font = Font(name="Calibri", bold=True, color=NAVY, size=11)
        ws.cell(r, 2).font = font_body
        ws.cell(r, 2).alignment = wrap
    ws["A10"] = "Reporting principles"
    ws["A10"].font = font_h2
    for i, (pid, text) in enumerate(parse_principles(), 11):
        ws.cell(i, 1, pid).font = Font(name="Calibri", bold=True, color=NAVY)
        ws.cell(i, 2, text).font = font_body
        ws.cell(i, 2).alignment = wrap
        ws.row_dimensions[i].height = 32
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100
    ws.row_dimensions[1].height = 32
    for r in range(3, 9):
        ws.row_dimensions[r].height = 22
    ws.print_title_rows = "1:1"


def index_sheet(wb: Workbook, sheets: list[tuple[str, str]]) -> None:
    ws = wb.create_sheet("Index", 1)
    add_title(ws, "Workbook index", "Each sheet is a catalog from the requirements pack.")
    headers = ["Sheet", "Contents"]
    rows = [[name, desc] for name, desc in sheets]
    write_table(ws, headers, rows, start_row=4)
    for r in range(5, 5 + len(rows)):
        cell = ws.cell(r, 1)
        name = cell.value
        cell.hyperlink = f"#'{name}'!A1"
        cell.font = Font(name="Calibri", color="0563C1", underline="single", size=11)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 88


def main() -> None:
    wb = Workbook()
    cover(wb)

    # Glossary
    ws = wb.create_sheet("Glossary")
    add_title(ws, "Glossary", "Canonical Operations language. Avoid-list is binding.")
    gloss = parse_glossary(ROOT / "CONTEXT.md")
    write_table(ws, ["Group", "Term", "Definition", "Avoid"], gloss, 4)

    # Scope / taxonomy from 01
    tax_headers, tax_rows = None, None
    for heading, header, rows in iter_tables(DOCS / "01-scope-and-taxonomy.md"):
        if header and header[0] == "ID" and "Subcategory" in header:
            ws = wb.create_sheet("Taxonomy")
            add_title(ws, "Operations subcategories (P01–P15)", heading)
            write_table(ws, header, rows, 4)
            tax_headers, tax_rows = header, rows
        elif header and header[0] == "Subject area":
            ws = wb.create_sheet("Scope")
            add_title(ws, "In / out of scope", "Adjacent subject areas and what Ops still sees.")
            write_table(ws, header, rows, 4)
        elif header and header[0] == "Persona":
            ws = wb.create_sheet("Personas")
            add_title(ws, "Personas and home dashboards")
            write_table(ws, header, rows, 4)

    # Pipeline phases — narrative fields + tables in 15
    phase_rows: list[list[str]] = []
    p_name = p_ms = p_proc = p_own = p_inn = p_out = ""
    def flush_phase():
        if p_name:
            phase_rows.append([p_name, p_ms, p_proc, p_own, p_inn, p_out])
    for line in (DOCS / "15-pipeline-phases.md").read_text(encoding="utf-8").splitlines():
        if line.startswith("## Phase"):
            flush_phase()
            p_name = line.lstrip("#").strip()
            p_ms = p_proc = p_own = p_inn = p_out = ""
        elif line.startswith("**Milestone:**"):
            p_ms = line.split(":", 1)[1].strip()
        elif line.startswith("**Process:**"):
            p_proc = line.split(":", 1)[1].strip()
        elif line.startswith("**Owner:**"):
            p_own = line.split(":", 1)[1].strip()
        elif line.startswith("**In:**"):
            p_inn = line.split(":", 1)[1].strip()
        elif line.startswith("**Out:**"):
            p_out = line.split(":", 1)[1].strip()
    flush_phase()
    if phase_rows:
        ws = wb.create_sheet("Pipeline_Phases")
        add_title(
            ws,
            "Pipeline phases",
            "Retail/Wholesale and Correspondent join at post-close. Two factories; do not average their clocks.",
        )
        write_table(ws, ["Phase", "Milestone", "Process", "Owner", "In", "Out"], phase_rows, 4)

    for heading, header, rows in iter_tables(DOCS / "15-pipeline-phases.md"):
        if header and header[0] == "Phase" and "Required extra slices" in " ".join(header):
            ws = wb.create_sheet("Phase_Dimensions")
            add_title(ws, "Extra dimensions by phase", "Always-on: Channel, product, purpose, team, Event or As-of date.")
            write_table(ws, header, rows, 4)
        elif header and header[0] == "Phase" and "Flow" in header:
            ws = wb.create_sheet("Phase_Metrics")
            add_title(ws, "Headline metrics by phase")
            write_table(ws, header, rows, 4)
        elif header and header[0] == "Control":
            ws = wb.create_sheet("Control_Overlays")
            add_title(ws, "Cross-cutting control (not a milestone)")
            write_table(ws, header, rows, 4)

    # Roles
    role_sheet_done = False
    action_rows: list[list[str]] = []
    for heading, header, rows in iter_tables(DOCS / "16-roles-and-actions.md"):
        if header[:2] == ["Role", "Typical home phase"] or (header and "Typical home phase" in header):
            ws = wb.create_sheet("Roles_Ops")
            add_title(ws, "Operations job roles")
            write_table(ws, header, rows, 4)
        elif header and header[0] == "Role" and "Who they are" in header:
            ws = wb.create_sheet("Roles_Counterparties")
            add_title(ws, "Counterparties (not Ops employees)")
            write_table(ws, header, rows, 4)
        elif header and header[0] == "Role" and "When they act" in header:
            ws = wb.create_sheet("Roles_Adjacent")
            add_title(ws, "Adjacent functions (handoff only)")
            write_table(ws, header, rows, 4)
        elif header == ["Role", "Actions"]:
            phase = heading
            for row in rows:
                action_rows.append([phase, row[0] if row else "", row[1] if len(row) > 1 else ""])
        elif header and header[0] == "Event" and "Who stamps" in header:
            ws = wb.create_sheet("Money_Stamps")
            add_title(ws, "Who may stamp money events")
            write_table(ws, header, rows, 4)
    if action_rows:
        ws = wb.create_sheet("Role_Actions")
        add_title(ws, "Role actions by pipeline phase", "Extracted from the Role | Actions tables in each phase.")
        write_table(ws, ["Phase", "Role", "Actions"], action_rows, 4)

    # Milestones and codes
    code_map = {
        "Canonical milestones": ("Milestones", "Canonical manufacturing milestones (MS-01–MS-13)"),
        "Waiting-on party codes": ("Waiting_On", "Four headline waiting-on codes. Blank is a data-quality fail."),
        "Fallout reason codes": ("Fallout_Codes", "Before Funded (R/W) or Correspondent purchase. Kickout is not Fallout."),
        "Kickout / investor suspense reasons": ("Kickout_Codes", "After Delivered. Map investor free text here."),
        "Manufacturing defect severity": ("Defect_Severity", "P12 headline is file-fail at critical."),
        "Condition class": ("Condition_Class", "PTD / PTF / PTP"),
        "Trailing document types (P10 heatmap)": ("Trailing_Docs", "Required set is product × investor."),
    }
    for heading, header, rows in iter_tables(DOCS / "11-milestones-and-reason-codes.md"):
        for key, (sheet, title) in code_map.items():
            if key.lower() in heading.lower() or heading.lower() in key.lower():
                ws = wb.create_sheet(sheet)
                add_title(ws, title)
                write_table(ws, header, rows, 4)
                break

    # Dimensions
    dim_rows: list[list[str]] = []
    for heading, header, rows in iter_tables(DOCS / "03-shared-dimensions.md"):
        # normalize to Group, Dimension, Definition, Grain / notes
        for row in rows:
            if len(row) >= 3:
                dim_rows.append([heading, row[0], row[1], row[2] if len(row) > 2 else "", row[3] if len(row) > 3 else ""])
    ws = wb.create_sheet("Dimensions")
    add_title(ws, "Shared (conformed) dimensions", "Channel is required on every enterprise view.")
    write_table(ws, ["Group", "Dimension", "Definition", "Grain / notes", "Do not confuse with"], dim_rows, 4)

    # Metrics + SLA from catalog
    sla_done = False
    metric_headers: list[str] = []
    metric_rows: list[list[str]] = []
    for heading, header, rows in iter_tables(DOCS / "04-metric-catalog.md"):
        if "ASSUMED SLA" in heading or header[:1] == ["Milestone / clock"]:
            ws = wb.create_sheet("SLA")
            add_title(ws, "ASSUMED SLA (not company policy)", "Initial LE send is TRID general business days, not calendar days.")
            write_table(ws, header, rows, 4)
            sla_done = True
            continue
        # skip tiny D01 headline lists if not tabular metrics
        if "ID" not in header and "id" not in [h.lower() for h in header]:
            continue
        if not metric_headers:
            metric_headers = ["Family"] + header
        # map into metric_headers
        idx = {h.lower(): i for i, h in enumerate(metric_headers)}
        for row in rows:
            rec = [""] * len(metric_headers)
            rec[0] = heading
            for h, v in zip(header, row):
                key = h.lower()
                if key not in idx:
                    metric_headers.append(h)
                    idx[key] = len(metric_headers) - 1
                    rec.append("")
                    rec[idx[key]] = v
                else:
                    rec[idx[key]] = v
            metric_rows.append(rec)
    # pad short rows
    width = len(metric_headers)
    metric_rows = [r + [""] * (width - len(r)) for r in metric_rows]
    ws = wb.create_sheet("Metrics")
    add_title(ws, "Metric catalog", "Business formulas, not SQL. IDs are stable.")
    write_table(ws, metric_headers, metric_rows, 4)

    # Events — union
    ev_h, ev_r = collect_tables(DOCS / "06-event-dictionary.md")
    ws = wb.create_sheet("Events")
    add_title(ws, "Event dictionary", "If an event is not stamped, the KPI is unavailable, not zero.")
    write_table(ws, ev_h, ev_r, 4)

    # Traceability
    for heading, header, rows in iter_tables(DOCS / "07-requirements-traceability.md"):
        if "Process" in header and "Primary metrics" in header:
            ws = wb.create_sheet("Traceability")
            add_title(ws, "Process → metrics → dashboards → events")
            write_table(ws, header, rows, 4)
        elif header and header[0] == "Risk":
            ws = wb.create_sheet("Counting_Rules")
            add_title(ws, "Counting collisions")
            write_table(ws, header, rows, 4)

    # Dashboards inventory
    for heading, header, rows in iter_tables(DOCS / "05-dashboard-inventory.md"):
        if header and header[0] == "ID" and "Dashboard" in header:
            ws = wb.create_sheet("Dashboards")
            add_title(ws, "Dashboard family", "D01 is the only enterprise home. D02 is the default volume/aging drill.")
            write_table(ws, header, rows, 4)

    # Dashboard specs — purpose + audience + KPIs as rows
    spec_rows = []
    for path in sorted((DOCS / "dashboards").glob("D*.md")):
        text = path.read_text(encoding="utf-8")
        did = path.stem.split("-")[0].upper()
        title = ""
        purpose = audience = cadence = kpis = filters = ""
        section = None
        buf: list[str] = []
        def take(sec, lines):
            return " ".join(x.strip() for x in lines if x.strip() and not x.startswith("#"))
        lines = text.splitlines()
        i = 0
        current = None
        blocks: dict[str, list[str]] = {}
        for line in lines:
            if line.startswith("# "):
                title = line[2:].strip()
            elif line.startswith("## "):
                current = line[3:].strip()
                blocks[current] = []
            elif current:
                blocks[current].append(line)
        purpose = take("p", blocks.get("Purpose and primary decision", [])[:8])
        audience = take("a", blocks.get("Audience and genre", [])[:6])
        cadence = take("c", blocks.get("Cadence and freshness", [])[:6])
        kpis = take("k", blocks.get("Headline KPIs", [])[:12])
        filters = take("f", blocks.get("Filters", [])[:8])
        spec_rows.append([did, title, purpose, audience, cadence, kpis, filters])
    ws = wb.create_sheet("Dashboard_Specs")
    add_title(ws, "Dashboard specifications (summary)", "Full layout/drill text remains in the markdown pack.")
    write_table(
        ws,
        ["ID", "Title", "Purpose / primary decision", "Audience and genre", "Cadence / freshness", "Headline KPIs", "Filters"],
        spec_rows,
        4,
    )

    # Stories + NFR
    stories = parse_stories(DOCS / "09-user-stories-and-nfr.md")
    ws = wb.create_sheet("User_Stories")
    add_title(ws, "User stories", "Acceptance is metric IDs and dashboard specs.")
    write_table(ws, ["ID", "Epic", "As a…", "I want…", "So that…"], stories, 4)
    for heading, header, rows in iter_tables(DOCS / "09-user-stories-and-nfr.md"):
        if header == ["ID", "Requirement"]:
            ws = wb.create_sheet("NFRs")
            add_title(ws, "Non-functional requirements")
            write_table(ws, header, rows, 4)

    # How to use: huddle, alerts, bands
    for heading, header, rows in iter_tables(DOCS / "12-how-to-use-the-numbers.md"):
        if "Minute" in header:
            ws = wb.create_sheet("Daily_Huddle")
            add_title(ws, "Daily huddle (30 minutes) — P15", "Do not use EXE in the daily huddle.")
            write_table(ws, header, rows, 4)
        elif header and header[0] == "Alert":
            ws = wb.create_sheet("Alerts")
            add_title(ws, "Recommended alerts (ASSUMED until Q1)")
            write_table(ws, header, rows, 4)
        elif header and "Rough non-bank" in " ".join(header):
            ws = wb.create_sheet("Industry_Bands")
            add_title(ws, "Industry-typical bands (not targets)")
            write_table(ws, header, rows, 4)

    # Open questions
    for heading, header, rows in iter_tables(DOCS / "08-phasing-and-open-questions.md"):
        if header and header[0] == "ID" and "Question" in header:
            ws = wb.create_sheet("Open_Questions")
            add_title(ws, "Open questions for Operations leadership", "Workshop Q1, Q3, Q4, Q6, Q9 first.")
            write_table(ws, header, rows, 4)

    # Anti-patterns
    for heading, header, rows in iter_tables(DOCS / "13-anti-patterns-and-validation.md"):
        if header and header[0] == "Anti-pattern":
            ws = wb.create_sheet("Anti_Patterns")
            add_title(ws, "Do not ship")
            write_table(ws, header, rows, 4)
        elif header and header[0] == "Check":
            ws = wb.create_sheet("Recon_Checks")
            add_title(ws, "Recon checks before calling a dashboard done")
            write_table(ws, header, rows, 4)

    # Personas already; wireframe map
    for heading, header, rows in iter_tables(DOCS / "wireframes" / "by-department.md"):
        if header and "Dept" in header[0] or (header and "Subcategory" in header[0]):
            ws = wb.create_sheet("Dept_Wireframes")
            add_title(ws, "Wireframes by department", "Interactive kit: docs/wireframes/index.html")
            write_table(ws, header, rows, 4)

    # Index last-created names
    skip = {"Cover", "Index"}
    catalog = []
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        title_cell = ws["A1"].value or ws.title
        catalog.append((ws.title, str(title_cell)))
    index_sheet(wb, catalog)

    # print settings
    for ws in wb.worksheets:
        ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
        ws.oddHeader.left.text = "Operations reporting requirements"
        ws.oddFooter.right.text = "Page &P of &N"

    wb.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size:,} bytes, {len(wb.worksheets)} sheets)")
    print("sheets:", [ws.title for ws in wb.worksheets])


if __name__ == "__main__":
    main()
